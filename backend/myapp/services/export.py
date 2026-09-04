from datetime import datetime
from pathlib import Path

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from myapp.models import ExportRecord

HEADER_FILL = PatternFill("solid", fgColor="E7F6F8")


def append_rows(worksheet, headers, rows):
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def create_excel_export(video, summary):
    output_dir = Path(settings.DOWNLOAD_ROOT)
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    identity = video.bvid or f"video_{video.pk}"
    file_name = f"analysis_{identity}_{timestamp}.xlsx"
    output_path = output_dir / file_name

    workbook = Workbook()
    overview = workbook.active
    overview.title = "视频概览"
    metadata = video.metadata or {}
    append_rows(
        overview,
        ["字段", "值"],
        [
            ["标题", video.title],
            ["BV号", video.bvid or ""],
            ["链接", video.source_url],
            ["UP主", video.owner_name],
            ["UP主ID", video.owner_id],
            ["时长（秒）", video.duration_seconds],
            ["播放量", metadata.get("views")],
            ["历史弹幕", metadata.get("historical_danmaku")],
            ["点赞", metadata.get("likes")],
            ["投币", metadata.get("coins")],
            ["收藏", metadata.get("favorites")],
            ["分享", metadata.get("shares")],
            ["本次分析弹幕数", summary.total_danmaku],
            ["每分钟平均弹幕", summary.metrics.get("average_per_minute")],
            ["峰值每秒弹幕", summary.metrics.get("peak_per_second")],
        ],
    )

    keywords = workbook.create_sheet("高频词")
    append_rows(
        keywords,
        ["排名", "关键词", "出现次数"],
        [
            [index, item.get("word"), item.get("count")]
            for index, item in enumerate(summary.top_words, start=1)
        ],
    )

    timeline = workbook.create_sheet("弹幕时间轴")
    append_rows(
        timeline,
        ["开始秒", "结束秒", "弹幕数"],
        [[item.get("start"), item.get("end"), item.get("count")] for item in summary.timeline],
    )

    peaks = workbook.create_sheet("精彩片段")
    append_rows(
        peaks,
        ["排名", "开始秒", "结束秒", "弹幕数"],
        [
            [index, item.get("start"), item.get("end"), item.get("count")]
            for index, item in enumerate(summary.peak_segments, start=1)
        ],
    )

    for worksheet in workbook.worksheets:
        for column in worksheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            worksheet.column_dimensions[column[0].column_letter].width = width

    workbook.save(output_path)
    file_url = settings.DOWNLOAD_URL + file_name
    ExportRecord.objects.create(
        video=video,
        file_name=file_name,
        file_url=file_url,
        format="xlsx",
    )
    return {"file_name": file_name, "file_url": file_url}
